from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

out = Path(r'D:\a22\docs\更新说明-项目更新至当前状态.docx')
out.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
doc.core_properties.title = 'A22 项目更新说明（拉取新项目后到当前状态）'

styles = doc.styles
styles['Normal'].font.name = '微软雅黑'
styles['Normal'].font.size = Pt(10.5)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('A22 项目更新说明')
run.bold = True
run.font.size = Pt(18)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('范围：从拉取新项目覆盖后，到当前本地已验证状态')

doc.add_paragraph('')
doc.add_paragraph('一、总体结论')
doc.add_paragraph('本次更新在“拉取并覆盖后的新项目”基础上完成了三类工作：恢复可运行配置、修复控件高/中/低风险边界问题、补齐前端校验并完成构建验证。最终本地状态为后端编译通过、前端构建通过。')

doc.add_paragraph('二、更新范围总览')
summary_rows = [
    ['backend/src/main/resources/application.yml', '配置更新', '后端数据库、Redis、Kafka、Flink 默认值与环境变量支持'],
    ['start-a22.ps1', '启动脚本更新', '默认账号、工具解析、前端/后端启动方式、端口与数据库参数传递'],
    ['backend/.../TaskExecutionEngine.java', '后端能力增强', 'concat / field-map / xml-json 等控件边界防护'],
    ['backend/.../MaskTransformPlugin.java', '后端能力增强', '短手机号/短身份证/短邮箱/短默认值脱敏保护'],
    ['backend/.../DataValidateTransformPlugin.java', '后端能力增强', 'required / 数值 / 日期 / allowed / regex 校验增强'],
    ['backend/.../DataRepairTransformPlugin.java', '后端能力增强', '维表读取 limit 上限告警'],
    ['frontend/src/FlowEditor.vue', '前端校验增强', 'filter / dedup / data-validate 的配置校验与提示'],
    ['.env / .env.cloud / docker-compose.cloud.yml', '环境与部署参考补齐', '本地启动变量及云部署参考配置'],
]
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr = table.rows[0].cells
hdr[0].text = '文件'
hdr[1].text = '更新类型'
hdr[2].text = '说明'
for r in summary_rows:
    row = table.add_row().cells
    row[0].text = r[0]
    row[1].text = r[1]
    row[2].text = r[2]

doc.add_paragraph('')
doc.add_paragraph('三、配置与启动相关更新')
items = [
    '1. 恢复后端默认数据库配置为本地可用账户（dev_lyb / lyb123456），并通过环境变量支持覆盖。',
    '2. 将 Redis / Kafka / Flink 默认连接恢复为 127.0.0.1，避免覆盖后仍指向远端地址导致连不上。',
    '3. 将 start-a22.ps1 默认 DbUsername 由 root@test 改为 dev_lyb，匹配当前 OceanBase 账户。',
    '4. 启动脚本增加工具自动解析（Maven/npm/mysql），兼容更多本地安装路径。',
    '5. 补齐本地 .env 与云部署参考文件（.env.cloud、docker-compose.cloud.yml）。'
]
for item in items:
    doc.add_paragraph(item, style='List Number')

doc.add_paragraph('四、后端控件修复汇总')
doc.add_paragraph('（一）高风险修复')
high = [
    'concat：null/“null”/“undefined”/空白统一规范化；模板未匹配占位符告警；结果过长截断保护。',
    'field-map：目标字段被多次映射时输出覆盖风险告警。',
    'xml-json：超长转换结果截断保护，避免下游写库失败。',
    'mask：短手机号、短身份证、短邮箱、短默认值全量脱敏，避免“几乎明文”。',
    'data-validate：required 扩展识别 “null/undefined/空白”；数值支持千分位/百分号/货币符号；日期支持多格式兜底；allowed 大小写不敏感。'
]
for item in high:
    doc.add_paragraph(item, style='List Bullet')

doc.add_paragraph('（二）中风险修复')
mid = [
    'concat：分隔符为空且存在多字段时告警，避免拼接结果不可逆。',
    'data-repair：维表读取命中 limit 上限时告警，避免误以为已完整修复。',
    'data-validate：空值命中后提前返回，避免后续无意义规则误报。'
]
for item in mid:
    doc.add_paragraph(item, style='List Bullet')

doc.add_paragraph('（三）低风险修复')
low = [
    'concat：target 与原记录字段重名时提示覆盖风险。',
    'data-validate：非法 regex 不再直接抛异常，改为记录“regex 配置异常”。'
]
for item in low:
    doc.add_paragraph(item, style='List Bullet')

doc.add_paragraph('五、前端校验更新')
fe = [
    'dedup：新增 orderFields 空值校验。',
    'filter：调整 value/date_between 校验顺序，避免跨字段过滤被误拦。',
    'data-validate：新增 rules 类型、空字段名、filterMode、issueField 长度校验。'
]
for item in fe:
    doc.add_paragraph(item, style='List Bullet')

doc.add_paragraph('六、本次验证结果')
ver = [
    '后端：mvn -DskipTests compile 通过。',
    '前端：npm run build 通过。',
    '建议回归样例：test-files/filter-test.csv、test-files/dedup-test.csv。'
]
for item in ver:
    doc.add_paragraph(item, style='List Bullet')

doc.add_paragraph('七、建议下一步')
nexts = [
    '1. 按“高-中-低”优先级回归核心控件场景。',
    '2. 将本次配置与启动脚本差异整理成团队统一开发环境说明。',
    '3. 若后续继续做自动化，可补充 filter/dedup/data-validate 的单测或接口测试。'
]
for item in nexts:
    doc.add_paragraph(item, style='List Number')

doc.save(out)
print(out)
