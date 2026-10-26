# -*- coding: utf-8 -*-
"""生成智流工坊演讲稿 Word 文档"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

doc = Document()

# ── 页面设置 ──
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.1)
section.right_margin = Inches(1.1)

# ── 样式 ──
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(12)
style.font.color.rgb = RGBColor(0x2C, 0x2C, 0x2C)
style.paragraph_format.space_after = Pt(8)
style.paragraph_format.line_spacing = 1.5
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

def set_run(p, text, size=12, bold=False, color='2C2C2C', italic=False, font='微软雅黑'):
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*bytes.fromhex(color))
    run.font.name = 'Calibri'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), font)
    return run

def add_doc_title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    p.paragraph_format.space_after = Pt(4)
    set_run(p, text, size=24, bold=True, color='1A5276')
    return p

def add_doc_subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    set_run(p, text, size=11, color='888888', italic=True)
    return p

def add_meta(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(24)
    set_run(p, text, size=10, color='999999')
    return p

def add_section_tag(text):
    """演讲段落标签，如【开场】"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(6)
    # 底部蓝色线条
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="8" w:space="4" w:color="2E74B5"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)
    set_run(p, text, size=14, bold=True, color='2E74B5')
    return p

def add_body(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.6
    set_run(p, text, size=12, color='333333')
    return p

def add_body_bold_prefix(prefix, rest):
    """带粗体前缀的段落"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.6
    set_run(p, prefix, size=12, bold=True, color='1A5276')
    set_run(p, rest, size=12, color='333333')
    return p

def add_bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.clear()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.5
    indent = 0.25 + level * 0.3
    p.paragraph_format.left_indent = Inches(indent)
    set_run(p, text, size=12, color='333333')
    return p

def add_bullet_bold_prefix(prefix, rest, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.clear()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.5
    indent = 0.25 + level * 0.3
    p.paragraph_format.left_indent = Inches(indent)
    set_run(p, prefix, size=12, bold=True, color='1A5276')
    set_run(p, rest, size=12, color='333333')
    return p

def add_quote(text):
    """引用/强调块"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.4)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="18" w:space="10" w:color="2E74B5"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)
    # 浅蓝背景
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EBF5FB"/>')
    pPr.append(shd)
    set_run(p, text, size=12, bold=True, color='1A5276')
    return p

def add_separator():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="4" w:space="1" w:color="D5D8DC"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)

# ================================================================
# 正文
# ================================================================

add_doc_title('智流工坊 · 项目介绍演讲稿')
add_doc_subtitle('Smart Flow Workshop — 实时数据流处理平台')
add_meta('演讲时长约 10-12 分钟　　版本：v1.0　　2026-08-16')

add_separator()

# ── 开场 ──
add_section_tag('【开场】')
add_body('大家好，今天给大家介绍我们团队开发的项目——智流工坊，英文名 Smart Flow Workshop。')
add_quote('一句话概括：这是一个实时数据流处理平台，让不懂代码的人也能通过拖拽画布，完成从数据采集、转换到入库的全流程。')

# ── 为什么要做 ──
add_section_tag('【为什么要做这个平台？】')
add_body('在实际的数据工程场景中，我们经常面对这样的问题：')
add_bullet('数据散落在 Excel、CSV、数据库、Kafka 消息队列等各种地方；')
add_bullet('每次做数据清洗、格式转换、脱敏处理，都要写一遍代码；')
add_bullet('流程改了，代码又要重写，维护成本很高。')
add_body('')
add_quote('智流工坊的目标就是：把数据管道的搭建从"写代码"变成"拖拖拽拽"。')

# ── 平台长什么样 ──
add_section_tag('【平台长什么样？】')
add_body('打开平台，登录后进入主界面，左侧是导航栏，包含这些核心页面：')
add_bullet_bold_prefix('概览', '——一眼看到所有任务的状态；')
add_bullet_bold_prefix('任务管理', '——创建、编辑、启停任务；')
add_bullet_bold_prefix('作业编排', '——最核心的功能，一个可视化的画布，从左侧面板拖出数据源节点、转换节点、输出节点，用线连起来，就完成了一条数据管道的编排；')
add_bullet_bold_prefix('数据工作区', '——上传文件、预览数据；')
add_bullet_bold_prefix('运行日志', '——查看任务执行的历史和状态；')
add_bullet_bold_prefix('Flink 作业', '——查看底层流处理引擎的运行情况。')
add_body('')
add_body('另外，平台内置了一个 AI 助手，你可以用自然语言告诉它需求，比如"读取 CSV，过滤金额大于 1000 的记录，脱敏后写入 MySQL"，它会自动生成作业编排方案。')

# ── 技术架构 ──
add_section_tag('【技术架构是怎么设计的？】')
add_body('整个平台分为三层：')
add_body('')

add_body_bold_prefix('第一层，前端交互层。', '')
add_body('基于 Vue 3 构建，用 Vite 做开发构建，核心的画布编排能力来自 Vue Flow 这个库。界面采用深色玻璃质感主题，科技感风格。')
add_body('')

add_body_bold_prefix('第二层，后端服务层。', '')
add_body('基于 Spring Boot 3.4，Java 17。负责业务逻辑、用户认证、文件解析、AI 集成等。这里重点说几个技术点：')
add_bullet('用 Spring Data JPA 做数据持久化，开发环境用 H2，生产用 MySQL；')
add_bullet('用 Spring Kafka 对接消息队列；')
add_bullet('用 Spring Data Redis 做缓存和 Token 管理；')
add_bullet('文件处理用了 Apache POI 读写 Excel，Apache Commons CSV 解析 CSV；')
add_bullet('AI 能力对接的是阿里云百炼的 DashScope API，默认用 qwen-plus 模型。')
add_body('')

add_body_bold_prefix('第三层，流处理引擎层。', '')
add_body('基于 Apache Flink。后端通过 Flink REST API 提交作业 JAR、查询状态和取消作业；Checkpoint、吞吐量与详细异常信息在 Flink Dashboard 中查看。')
add_body('')
add_body('基础设施通过 Docker Compose 一键编排，包括 MySQL 8.4、Redis 7、Kafka 3.8、Flink 1.20，开箱即用。')

# ── 转换插件体系 ──
add_section_tag('【核心能力：转换插件体系】')
add_body('平台的转换能力是插件化的。我们定义了一个 TransformPlugin 接口，每个转换算子都是一个独立插件，可以自由组合。目前已实现的插件包括：')
add_bullet_bold_prefix('字段映射', '——重命名字段；')
add_bullet_bold_prefix('条件过滤', '——按表达式筛选数据；')
add_bullet_bold_prefix('去重', '——按指定字段去除重复记录；')
add_bullet_bold_prefix('类型转换', '——字符串转数字、日期格式化；')
add_bullet_bold_prefix('数据脱敏', '——手机号、身份证号遮罩处理；')
add_bullet_bold_prefix('字段拼接', '——多个字段合并；')
add_bullet_bold_prefix('XML/JSON 互转', '——格式适配。')
add_body('')
add_body('后续要扩展新的转换能力，只需要实现接口、注册到插件仓库即可，不需要改动引擎代码。')

# ── 数据源和目标 ──
add_section_tag('【支持的数据源和目标】')
add_body('数据源方面，目前支持：JSON 文件、CSV 文件、Excel 文件、MySQL 表、Kafka 消息。')
add_body('数据目标方面，支持写入 MySQL、Redis、Kafka，以及导出为 CSV/Excel 文件。')
add_body('基本上覆盖了常见的数据工程场景。')

# ── 权限管理 ──
add_section_tag('【权限管理】')
add_body('平台实现了基于角色的访问控制，四种角色：超级管理员、管理员、普通用户、运维只读。不同角色看到的页面不同，操作权限也不同。比如运维只读角色只能查看，不能修改任务配置。')

# ── 怎么跑起来 ──
add_section_tag('【怎么跑起来？】')
add_body('开发环境搭建非常简单：')
add_body('')
add_body_bold_prefix('第一步，', 'docker-compose up -d ——拉起 MySQL、Redis、Kafka、Flink；')
add_body_bold_prefix('第二步，', 'cd backend && mvn spring-boot:run ——启动后端；')
add_body_bold_prefix('第三步，', 'cd frontend && npm ci && npm run dev ——启动前端；')
add_body_bold_prefix('第四步，', '打开 http://localhost:5173 就可以用了。')

# ── 总结 ──
add_section_tag('【总结】')
add_body('最后总结一下智流工坊的核心价值：')
add_body('')
add_quote('低门槛——拖拽编排，不写代码；\n全链路——采集、转换、存储、监控一条龙；\n智能化——AI 辅助生成数据管道；\n可扩展——插件架构，按需扩展转换能力。')
add_body('')
add_body('这就是智流工坊。谢谢大家，欢迎提问。')

# ── 页脚 ──
add_separator()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(8)
set_run(p, '— 智流工坊 · Smart Flow Workshop —', size=9, color='999999')

# ── 保存 ──
out = r'D:\a22\docs\智流工坊-演讲稿.docx'
doc.save(out)
print(f'OK: {out}')
