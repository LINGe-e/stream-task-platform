from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from pathlib import Path

output_path = Path('D:/a22/docs/智流工坊-项目演讲稿-2026-09-14.docx')

content_sections = [
    {
        'title': '开场：一句话说清楚这个项目在什么场景下解决什么问题',
        'body': [
            '各位好，今天我介绍的项目是“智流工坊（Smart Flow Workshop）”。它主要应用于企业内部的数据治理、跨系统数据集成和实时流处理场景，用来帮助业务团队把“分散在文件、数据库和消息队列里的数据”清洗、转换、对齐，再稳定地写入目标系统。',
            '简单来说，它的价值就是：让不懂代码的业务同事，也能通过拖拽的方式，把一条数据流从源头编排到目的地，同时满足校验、脱敏、补全和监控需求。'
        ]
    },
    {
        'title': '第一步：从一个真实业务痛点讲起',
        'body': [
            '在很多项目里，数据并不是一开始就是“干净”的。现实中最常见的问题有三类：第一，字段不统一；第二，敏感信息需要保护；第三，数据来源太杂，人工处理容易出错。',
            '智流工坊就是围绕这些问题设计的。我们不是去写一次性脚本，而是把处理流程变成“可编排、可复用、可监控”的任务。'
        ]
    },
    {
        'title': '第二步：用什么场景来落地',
        'body': [
            '第一类场景，是日常数据清洗与入库。比如运营上传 CSV、Excel 或 JSON 文件，需要统一字段名、过滤异常值、去除重复记录，最后写入 MySQL。平台里提供 JSON 输入、CSV 输入、Excel 输入、MySQL 输出等控件，可以直接拖拽组装。',
            '第二类场景，是跨源数据打通。比如需要把 JDBC 外部表、Redis、Kafka 里的数据统一归集，再按规则映射到目标表。平台里的 JDBC 输入、Redis 输入、Kafka 输入、字段映射、Redis 查询补全，就是为这种场景设计的。',
            '第三类场景，是数据质量与合规。很多团队在做业务数据时，不仅要处理数据，还要检查手机号、金额、状态、日期是否合法，同时对敏感字段做脱敏处理。平台内置了数据校验、字段脱敏、数据补全等算子，能够直接在流程里完成这些问题处理。'
        ]
    },
    {
        'title': '第三步：把项目功能一点一点讲清楚',
        'body': [
            '在平台里，核心操作路径是：任务管理 → 可视化编排 → 节点测试 → 执行与监控。',
            '首先是任务管理。平台支持任务创建、编辑、发布、撤回、运行和查看运行记录。任务可以设置为“一次执行”或“持续处理”，还能配置固定间隔、每日定时、Cron 表达式等调度方式。',
            '其次是可视化编排。前端基于 Vue 3 和 FlowEditor 实现了拖拽式画布，左侧是输入、处理、输出和控制控件，右侧可以配置每个节点的参数。系统会自动校验 DAG 合法性，避免循环连线或非法编排。',
            '然后是数据处理能力。平台目前支持字段映射、条件过滤、去重、字段拼接、类型转换、字段脱敏、数据补全、Redis 查询补全、数据校验、XML/JSON 转换等处理节点，基本覆盖了常见的数据治理需求。',
            '最后是输出和执行。处理结果可以写入 MySQL、Kafka、JDBC 外部表、CSV 或 Excel，同时支持导出文件下载。运行模式包括一次执行和持续处理，任务状态会在草稿、已发布、运行中、成功、失败之间流转，方便运营和运维同事一起协作。'
        ]
    },
    {
        'title': '第四步：重点突出项目里的特色',
        'body': [
            '第一个特色，是“低门槛的可视化编排”。很多数据平台功能很强，但学习成本太高。智流工坊强调的是业务同事也能上手，用拖拽和连线表达数据流，不需要一开始就理解复杂代码。',
            '第二个特色，是“单节点测试能力”。在前端 FlowEditor 中，用户可以单独对某个处理节点进行测试，输入样本数据后立即查看输出结果。这对排查字段映射、过滤条件、脱敏规则是否正确非常有帮助。',
            '第三个特色，是“AI 辅助编排”。平台内置 AI 助手，用户可以用自然语言描述需求，例如“读取 CSV，过滤金额大于 1000，字段脱敏，写入 MySQL”，系统会结合当前页面上下文给出流程建议。这里的设计重点不是让 AI 直接操作数据库，而是让它帮助用户更快理解该怎么编排。',
            '第四个特色，是“任务运行模式完整”。平台不仅支持一次执行，还支持持续处理、重试机制、增量状态和增量游标。这意味着它不仅是一个“跑一次的工具”，而是可以承担更接近生产环境的长期数据任务。',
            '第五个特色，是“权限和协作机制”。平台支持高级管理员、管理员、普通用户和只读运维四种角色，不同角色拥有不同的创建、编辑、发布和执行权限。这样既方便业务同事使用，也方便管理员统一管理。'
        ]
    },
    {
        'title': '第五步：用一句话总结项目价值',
        'body': [
            '总的来说，智流工坊不是单纯做一个“文件导入工具”，而是希望成为一个“轻量但完整”的数据流处理平台。它覆盖了从数据接入、数据转换、质量校验，到任务编排、调度执行和运行监控的完整链路。',
            '如果用一句话总结：它是面向数据团队的可视化流处理工作台，帮助企业把复杂的数据处理流程变得可配置、可复用、可监控。谢谢大家。'
        ]
    }
]

doc = Document()
p = doc.add_paragraph()
p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run = p.add_run('智流工坊项目演讲稿')
run.bold = True
run.font.size = Pt(20)

p = doc.add_paragraph()
p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
run = p.add_run('应用场景介绍 · 功能逐步展开 · 项目特色突出')
run.font.size = Pt(11)

for section in content_sections:
    doc.add_paragraph(section['title'], style='Heading 1')
    for text in section['body']:
        doc.add_paragraph(text, style='List Bullet')

doc.add_paragraph('说明', style='Heading 2')
doc.add_paragraph('本演讲稿基于当前仓库中的 README、前端 FlowEditor、后端 TaskService/TaskExecutionEngine 以及现有项目文档整理，突出项目真实功能与演示亮点，适合在项目汇报、演示评审、方案介绍时直接使用。', style='List Bullet')

doc.core_properties.title = '智流工坊项目演讲稿'
doc.core_properties.subject = '智流工坊应用场景与项目特色介绍'
doc.core_properties.comments = '根据当前项目源码与文档整理的演讲稿素材。'

output_path.parent.mkdir(parents=True, exist_ok=True)
doc.save(output_path)
print(output_path)
