from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path('docs') / '智流工坊-项目可行性分析与成本模型.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
BLACK = RGBColor(0, 0, 0)

def set_run(run, name='宋体', size=12, bold=False):
    run.font.name = name; run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size); run.font.bold = bold; run.font.italic = False; run.font.color.rgb = BLACK

def set_style(style, name='宋体', size=12, bold=False):
    style.font.name = name; style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    style.font.size = Pt(size); style.font.bold = bold; style.font.color.rgb = BLACK

def set_cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr(); borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None: borders = OxmlElement('w:tcBorders'); tcPr.append(borders)
    for edge in ('top','left','bottom','right'):
        tag = 'w:' + edge; el = borders.find(qn(tag))
        if el is None: el = OxmlElement(tag); borders.append(el)
        for k,v in [('val','single'),('sz','8'),('space','0'),('color','000000')]: el.set(qn('w:'+k), v)

def set_cell(cell, text, bold=False, center=False):
    cell.text = ''; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_border(cell)
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.15
    r = p.add_run(str(text)); set_run(r, '宋体', 10.5, bold)
    tcPr = cell._tc.get_or_add_tcPr(); shd = tcPr.first_child_found_in('w:shd')
    if shd is None: shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),'FFFFFF'); shd.set(qn('w:val'),'clear')

def add_table(doc, rows, widths, caption):
    table = doc.add_table(rows=len(rows), cols=len(widths)); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    for i,row in enumerate(rows):
        for j,value in enumerate(row): table.cell(i,j).width = Cm(widths[j]); set_cell(table.cell(i,j), value, i == 0, i == 0)
    p = doc.add_paragraph(style='题注'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(8); set_run(p.add_run(caption),'宋体',10.5)

def body(doc, text, indent=True):
    p=doc.add_paragraph(style='正文'); p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.ONE_POINT_FIVE; p.paragraph_format.space_after=Pt(6)
    if indent: p.paragraph_format.first_line_indent=Cm(0.74)
    set_run(p.add_run(text),'宋体',12); return p

def heading(doc, text, level):
    p=doc.add_paragraph(style={1:'标题1',2:'标题2',3:'标题3'}[level]); p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_before=Pt(12 if level==1 else 8); p.paragraph_format.space_after=Pt(4)
    set_run(p.add_run(text),'黑体',{1:16,2:14,3:12}[level],True); return p

def label(doc, text):
    p=doc.add_paragraph(style='小标题'); p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.ONE_POINT_FIVE; p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(2); set_run(p.add_run(text),'黑体',12,True)

def numbered(doc, items):
    for item in items: body(doc, item)

def build():
    doc=Document(); sec=doc.sections[0]; sec.top_margin=sec.bottom_margin=Cm(2.54); sec.left_margin=sec.right_margin=Cm(3.18)
    styles=doc.styles; set_style(styles['Normal'],'宋体',12)
    for name,fn,sz,bold in [('正文','宋体',12,False),('小标题','黑体',12,True),('标题1','黑体',16,True),('标题2','黑体',14,True),('标题3','黑体',12,True),('题注','宋体',10.5,False)]:
        try: st=styles[name]
        except KeyError: st=styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        st.base_style=styles['Normal']; set_style(st,fn,sz,bold)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(30); p.paragraph_format.space_after=Pt(12); set_run(p.add_run('智流工坊'),'黑体',20,True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(24); set_run(p.add_run('项目可行性分析与成本模型'),'黑体',16,True)

    heading(doc,'第1章 引言',1); heading(doc,'1.1 项目简介',2)
    numbered(doc,['（1）软件名称：智流工坊','（2）项目定位：面向数据工程团队的可视化实时数据流处理平台','（3）承办团队：智流工坊项目组'])
    heading(doc,'1.2 项目背景',2)
    body(doc,'在数据来源日益多样化、业务处理时效要求不断提高的背景下，企业通常需要同时面对文件、数据库、消息队列和对象存储等多种数据源。传统数据管道依赖人工编写脚本和配置，存在开发门槛高、流程复用难、运行状态不透明以及故障定位成本高等问题。')
    body(doc,'智流工坊以可视化任务编排为核心，将数据接入、字段处理、实时计算、结果写入和运行监控整合到统一工作台中。平台当前已具备文件拖拽上传、JSON/CSV/Excel 解析、数据预览、任务管理、运行记录和基础 Flink 集群状态查看能力，并已支持 Kafka 与外部 JDBC 数据库链路；Redis、HDFS、S3 和 CDC 等链路可继续扩展，目标是降低数据流处理的使用门槛，提升数据开发与运维效率。')

    heading(doc,'第2章 成本模型',1); heading(doc,'2.1 软件购买成本',2)
    body(doc,'智流工坊采用前后端分离和开源组件优先的技术路线，软件成本主要来自部署环境、域名、证书、日志与可选的托管服务。以下金额为项目初期预算估算，实际费用需根据部署规模和服务商报价确认。')
    numbered(doc,['（1）前端开发框架：Vue 3、Vite、Vue Flow 等开源组件，不产生授权购买费用。','（2）后端开发框架：Spring Boot、Maven 和 Java 17，采用开源许可，不产生授权购买费用。','（3）数据基础设施：MySQL、Redis、Kafka、Apache Flink 等可使用开源版本；若采用云托管实例，则按实际资源量计费。','（4）部署与辅助服务：域名、HTTPS 证书、对象存储、日志和备份等，按项目初期预算约 5000 元。'])
    add_table(doc,[['系统开发所需要的软件/服务','费用（元）'],['开源框架与数据库组件','0'],['部署、域名、证书及辅助服务（预算）','5000'],['软件成本总计（预算）','5000']],[10.5,5.5],'表2-1 软件成本')
    heading(doc,'2.2 硬件购买成本',2)
    body(doc,'智流工坊的开发和验证需要前端、后端、数据库及 Flink 集群等运行环境。为保证本地联调和小规模演示的稳定性，建议配置开发终端、测试服务器和备份存储；生产环境的硬件或云资源应结合并发量、数据量和保留周期另行核算。')
    numbered(doc,['（1）开发设备：开发与测试终端 3 台，按每台约 10000 元估算，共计 30000 元。','（2）测试与演示服务器：用于 MySQL、Redis、Kafka、Flink 等组件联调，按一年资源费用约 36000 元估算。','（3）备份和网络设备：用于样例数据、任务配置及运行日志的备份，预算约 10000 元。'])
    add_table(doc,[['系统开发所需的硬件/资源','费用（元）'],['开发与测试终端（3 台）','30000'],['测试与演示服务器（1 年预算）','36000'],['备份与网络设备（预算）','10000'],['硬件及资源成本总计（预算）','76000']],[10.5,5.5],'表2-2 硬件及资源成本')
    heading(doc,'2.3 系统开发费用',2); heading(doc,'2.3.1 人员配备',3)
    body(doc,'智流工坊的研发工作覆盖产品设计、前端交互、后端接口、数据处理引擎、测试和部署运维等环节，建议配置如下角色：')
    add_table(doc,[['人员配备','数量（人）','职能'],['项目负责人/产品经理','1','统筹需求、进度、验收和协作管理'],['前端开发工程师','1','负责 Vue 管理台、可视化画布和交互体验'],['后端开发工程师','2','负责 Spring Boot API、认证、任务和数据服务'],['数据工程/Flink 工程师','1','负责数据源、转换算子、Flink 作业和运行链路'],['测试与运维工程师','1','负责测试验证、部署、监控和故障排查']],[6.0,2.5,7.5],'表2-3-1 人员配备表')
    body(doc,'本项目按 6 个月研发周期进行预算估算，实际人员投入可根据迭代节奏和交付范围调整。')
    add_table(doc,[['系统开发人员配备','月薪（元）','时间（月）','研发过程薪酬（元）'],['项目负责人/产品经理','18000','6','108000'],['前端开发工程师','16000','6','96000'],['后端开发工程师','18000','6','216000（2 人）'],['数据工程/Flink 工程师','20000','6','120000'],['测试与运维工程师','14000','6','84000'],['系统开发薪酬成本（预算）','','','624000']],[6.0,3.0,2.5,4.5],'表2-3 项目薪酬成本表')
    heading(doc,'2.4 总体成本估计',2)
    body(doc,'除软件、硬件及人员成本外，还需要考虑办公协作、版本发布、测试数据维护、备份、培训和不可预见支出等费用，按 20000 元进行初步预算。根据当前估算，智流工坊首期 6 个月研发与演示环境的总投入约为 725000 元；该金额为项目预算模型，不等同于已经发生的采购或薪酬支出。')
    add_table(doc,[['成本类型','费用（元）'],['软件及辅助服务成本（预算）','5000'],['硬件及资源成本（预算）','76000'],['系统开发薪酬成本（预算）','624000'],['其他费用（预算）','20000'],['首期预算总费用','725000']],[10.5,5.5],'表2-4 总成本表')

    heading(doc,'第3章 可行性分析',1); heading(doc,'3.1 风险评估及策略',2)
    body(doc,'智流工坊涉及数据接入、任务编排、实时计算、外部组件连接和权限管理等多个环节，开发与运营过程中可能面临技术、运行和数据安全风险。针对主要风险，项目应在迭代过程中同步建立验证、监控和应急机制。')
    label(doc,'技术风险')
    add_table(doc,[['风险特征','策略'],['不同数据源的协议、字段类型和网络环境存在差异，可能导致任务编排成功但运行时连接或数据解析失败。','1、为数据源、转换算子和输出目标定义统一接口及参数校验。\n2、在任务发布前执行图结构、字段和连接配置校验。\n3、通过节点测试、样例数据和端到端演练提前发现问题。\n4、对 Kafka、Flink、数据库等外部依赖增加超时、重试和可观测日志。'],['任务规模增长后，Flink 并行度、消息积压、数据库写入和任务资源可能成为性能瓶颈。','1、在上线前进行负载和压力测试，记录吞吐量、延迟和失败率。\n2、根据数据规模调整并行度、批量写入和资源配额。\n3、对任务运行状态、消费积压、Slot 使用和异常日志建立监控。']],[8.0,8.0],'表2-5 技术风险')
    label(doc,'运营风险')
    add_table(doc,[['风险特征','策略'],['用户配置的任务数量和数据量持续增长，现有测试环境资源不足，导致页面响应变慢或作业提交失败。','1、区分开发、测试和生产资源，设置任务与文件大小上限。\n2、定期清理临时数据和历史运行记录，建立备份与恢复策略。\n3、对关键服务设置健康检查、资源告警和故障处理流程。'],['平台功能较多，用户对节点参数、任务状态和失败原因理解不足，影响使用效率。','1、提供节点说明、配置示例和错误信息。\n2、保留任务运行日志与结果预览，支持按运行记录定位问题。\n3、通过反馈收集和版本迭代持续优化交互和帮助内容。']],[8.0,8.0],'表2-6 运营风险')
    label(doc,'法律与隐私风险')
    add_table(doc,[['风险特征','策略'],['上传文件、账号信息、任务配置和运行数据可能包含业务敏感内容，若权限或传输保护不足，可能造成数据泄露。','1、遵循最小权限原则，完善登录认证、角色权限和任务协作边界。\n2、对传输链路、密码、令牌和敏感配置进行安全保护，避免将真实凭据写入代码或文档。\n3、对上传文件、运行日志和导出结果设置访问控制、保留周期及审计记录。\n4、定期进行依赖升级、漏洞扫描和安全检查。']],[8.0,8.0],'表2-7 法律与隐私风险')

    heading(doc,'3.2 技术可行性',2); heading(doc,'3.2.1 技术力量可行性',3)
    numbered(doc,['（1）项目采用 Vue 3、Spring Boot、Java 17、MySQL、Redis、Kafka 和 Apache Flink 等成熟技术，社区资料和工程实践较为丰富。','（2）代码按前端、后端、引擎、任务、文件、认证和 Flink 对接等模块组织，便于由具备 Web 开发、数据工程和运维能力的成员分工实施。','（3）项目已有文件解析、任务运行、数据转换插件、Flink 提交和状态查询等实现基础，可在现有工程上持续迭代。'])
    heading(doc,'3.2.2 技术方案可行性',3)
    numbered(doc,['（1）以可视化 DAG 作为任务表达方式，通过节点、连线和参数配置描述数据处理流程，并在发布前进行图结构校验。','（2）后端提供统一 API 和任务执行引擎，能够承接文件、JDBC、Kafka、Redis、HDFS 和 S3 等数据源或目标的扩展。','（3）转换插件采用统一注册与执行机制，支持字段映射、过滤、去重、脱敏、数据修复、类型转换、拼接和 XML/JSON 转换等处理。','（4）Flink 作为实时执行扩展，通过作业提交、集群健康检查和运行状态查看形成可观测的执行链路；在本地演示场景下也可使用后端执行能力进行验证。'])

    heading(doc,'3.3 经济可行性',2); heading(doc,'3.3.1 成本预算',3)
    body(doc,'项目优先采用开源框架和可替换的基础设施组件，软件授权成本较低；主要投入集中在研发人员、测试环境和运行资源。通过分阶段建设，先完成文件接入、任务编排和基础执行闭环，再按实际用户量扩充 Flink、Kafka、存储和监控资源，可以控制前期投入并降低一次性采购风险。')
    heading(doc,'3.3.2 预期收益',3)
    body(doc,'智流工坊的直接收益包括减少重复脚本开发、缩短数据任务交付周期、降低任务运维和故障定位成本，并提升数据处理流程的复用程度。对于数据工程团队，平台还可沉淀标准化节点、任务模板和运行经验，为后续提供内部数据服务、项目交付或平台化能力输出创造条件。')
    heading(doc,'3.3.3 经济可行性总结',3)
    body(doc,'在采用开源技术、分阶段投入和按需扩容的前提下，智流工坊的成本结构清晰，研发投入可通过效率提升和任务复用逐步回收，具有较好的经济可行性。预算中的硬件、云资源和人员单价应在立项或采购阶段根据实际报价重新确认。')

    heading(doc,'3.4 应用环境可行性',2); heading(doc,'3.4.1 市场需求',3)
    body(doc,'企业数据来源和实时分析场景持续增加，数据团队需要更加直观、可复用和可观测的数据处理工具。智流工坊面向文件导入、数据清洗、任务编排和实时流处理等常见场景，能够覆盖从样例验证到持续运行的基础需求。')
    heading(doc,'3.4.2 技术与部署环境',3)
    body(doc,'项目支持在具备 JDK 17、Maven、Node.js 20、Docker Desktop 和 MySQL 等条件的环境中进行开发与演示，并可通过 Docker Compose 组织 MySQL、Redis、Kafka 和 Flink 等本地基础服务。前后端分离结构和配置化数据库连接便于迁移到测试或生产环境。')
    heading(doc,'3.4.3 法律与政策环境',3)
    body(doc,'平台应根据部署行业和数据类型落实网络安全、数据安全、个人信息保护及开源软件许可等要求。对于上传文件、用户账号、任务配置和运行日志，应明确数据用途、访问范围、留存周期与责任边界；在涉及敏感或个人数据的场景中，完成必要的脱敏、加密、审计和合规评估。')

    heading(doc,'3.5 产品推广可行性',2); heading(doc,'3.5.1 产品（Product）与服务（Service）',3)
    body(doc,'产品策略：'); numbered(doc,['（1）差异化定位：以可视化 DAG 编排和实时数据流处理为核心，将数据接入、转换、运行和结果查看集中在同一工作台。','（2）体验优化：通过拖拽上传、JSON 预览、节点参数配置、任务状态卡片、运行日志和失败分析降低使用门槛。','（3）能力扩展：持续增加数据源、处理插件、结果目标、任务模板和监控告警能力，使平台覆盖更多数据工程场景。'])
    body(doc,'服务策略：'); numbered(doc,['（1）提供安装部署、环境检查、数据接入和任务编排示例，帮助用户快速完成首次任务。','（2）建立问题反馈、运行日志和版本发布机制，持续收集任务失败、性能和体验问题。','（3）针对企业使用场景提供权限配置、任务协作、备份恢复、升级维护和安全检查等配套服务。'])
    heading(doc,'3.5.2 渠道（Channel）',3); body(doc,'线上渠道：')
    numbered(doc,['（1）通过项目文档、演示环境和示例数据展示文件接入、可视化编排与实时运行效果。','（2）在开发者社区、技术交流活动和团队内部知识库发布使用指南、节点案例和实践经验。','（3）通过试用反馈和典型任务模板形成可复制的行业解决方案，扩大平台使用范围。'])
    body(doc,'线下渠道：'); numbered(doc,['（1）面向企业数据团队、研发团队和高校实验室开展现场演示与培训。','（2）与数据咨询、系统集成和云服务团队合作，结合具体项目导入平台能力。'])

    heading(doc,'3.6 结论和总结',2)
    body(doc,'综上所述，智流工坊在技术、经济和应用环境方面具备较好的可行性。项目已经形成前端管理台、后端 API、数据文件解析、任务编排、运行记录和 Flink 对接的工程基础，能够支撑从文件导入到任务执行与结果查看的核心闭环。采用开源组件和分阶段建设策略，可将前期成本控制在可规划范围内。后续实施应重点关注真实数据源接入、Flink 集群稳定性、权限与数据安全、任务监控告警以及大规模数据下的性能验证，并根据实际用户反馈持续调整产品边界和资源投入。')

    for p in doc.paragraphs:
        for r in p.runs:
            r.font.color.rgb=BLACK
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb=BLACK
    doc.save(OUT); print(OUT)

if __name__=='__main__': build()

