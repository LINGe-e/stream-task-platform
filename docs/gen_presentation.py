# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

doc = Document()

style = doc.styles['Normal']
style.font.name = '微软雅黑'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
style.paragraph_format.line_spacing = 1.5

for level, size, color in [(1, 22, '1A3667'), (2, 16, '2E5D8A'), (3, 13, '3A7BB8')]:
    h = doc.styles[f'Heading {level}']
    h.font.name = '微软雅黑'
    h.font.size = Pt(size)
    h.font.color.rgb = RGBColor.from_string(color)
    h.font.bold = True
    h.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

def add_para(text, bold=False, size=None, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = '微软雅黑'
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    if bold: run.font.bold = True
    if size: run.font.size = Pt(size)
    if color: run.font.color.rgb = RGBColor.from_string(color)
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def add_bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.clear()
    run = p.add_run(text)
    run.font.name = '微软雅黑'
    run.font.size = Pt(11)
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    if level > 0: p.paragraph_format.left_indent = Cm(level * 0.75)
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(10)
                run.font.name = '微软雅黑'
                run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)
                    run.font.name = '微软雅黑'
                    run.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    doc.add_paragraph()
    return table

# === 封面 ===
for _ in range(4): doc.add_paragraph()
add_para('A22 智流工坊', bold=True, size=36, color='1A3667', align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
add_para('通用流处理任务管理平台', bold=True, size=20, color='2E5D8A', align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)
add_para('项目讲解流程稿', bold=False, size=14, color='666666', align=WD_ALIGN_PARAGRAPH.CENTER, space_after=36)
add_para('主讲人：凌一搏', size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
add_para('配合：徐晋辉', size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
add_para('版本：v2.0 | 2026年9月', size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
doc.add_page_break()

# === 目录 ===
doc.add_heading('目 录', level=1)
add_para('')
toc = ['一、开场与项目定位','二、应用场景与价值','三、系统架构总览','四、核心功能详解',
'    4.1 数据工作区','    4.2 作业编排 - 可视化 DAG 流程设计器','    4.3 转换处理插件体系',
'    4.4 多种输出目标','    4.5 任务管理与调度','    4.6 持续处理与增量模式',
'    4.7 AI 智能助手','    4.8 Flink 作业集成','    4.9 用户与权限管理','    4.10 运行日志与监控',
'五、技术栈与部署方案','六、演示流程建议','七、常见问题与答疑要点','八、总结与展望']
for item in toc: add_para(item, size=12, space_after=4)
doc.add_page_break()

# === 一 ===
doc.add_heading('一、开场与项目定位', level=1)
doc.add_heading('1.1 开场白', level=2)
add_para('"大家好，今天给大家介绍我们的项目 —— A22 智流工坊，一个通用流处理任务管理平台。"')
doc.add_heading('1.2 痛点引入', level=2)
add_para('在日常数据处理工作中，常见问题：')
add_bullet('数据来源分散：文件（JSON/CSV/Excel）、数据库（MySQL）、消息队列（Kafka）、缓存（Redis）')
add_bullet('处理流程复杂：字段映射、清洗、脱敏、去重等需反复手动操作')
add_bullet('缺乏可视化：处理流程是"黑盒"，出问题难以定位')
add_bullet('协作困难：任务配置散落在脚本中，团队难以复用')
add_bullet('监控缺失：任务运行状态不透明')
doc.add_heading('1.3 项目定位', level=2)
add_para('智流工坊定位为"一站式流数据处理平台"，核心理念：')
add_bullet('让非技术人员也能通过拖拽完成数据处理流程设计')
add_bullet('让技术人员获得更强的编排能力和扩展性')
add_bullet('让团队协作更高效，任务管理更规范')
doc.add_page_break()

# === 二 ===
doc.add_heading('二、应用场景与价值', level=1)
doc.add_heading('2.1 典型应用场景', level=2)
scenarios = [
    ('场景一：多源数据整合', '企业需将 CRM 导出的 CSV、ERP 的 JSON 接口、业务数据库的订单表统一汇总。传统方式需编写多个脚本，使用智流工坊只需拖拽配置即可完成。'),
    ('场景二：数据清洗与脱敏', '原始数据含空值、异常值、重复记录，且涉及隐私字段需脱敏。平台内置条件过滤、去重、字段脱敏、数据校验等插件，无需编写代码。'),
    ('场景三：实时数据流处理', 'Kafka 消费 → 数据转换 → 写入数据库等持续运行场景，支持增量消费模式和 Redis 位点记录，7x24 小时运行。'),
    ('场景四：跨系统数据同步', 'MySQL 订单数据同步到 Redis 缓存，或 JDBC 数据源写入 Kafka，支持多种数据源自由组合。'),
    ('场景五：报表数据准备', '从数据库提取数据、汇总计算、输出为 Excel 报表，配合定时调度实现自动化。'),
]
for title, desc in scenarios:
    add_para(title, bold=True, size=12, color='2E5D8A')
    add_para(desc)

doc.add_heading('2.2 项目价值对比', level=2)
add_table(['维度', '传统方式', '智流工坊'], [
    ['上手门槛', '需编写代码', '可视化拖拽，零代码'],
    ['开发效率', '一个流程可能需数天', '数小时内完成'],
    ['可维护性', '脚本分散，难复用', '统一管理，版本可追溯'],
    ['协作能力', '依赖口头沟通', '任务可视化，团队共享'],
    ['监控能力', '日志分散，排查困难', '集中监控，实时状态'],
    ['扩展性', '改动需重新编码', '插件化架构，灵活扩展'],
])
doc.add_page_break()

# === 三 ===
doc.add_heading('三、系统架构总览', level=1)
doc.add_heading('3.1 整体架构', level=2)
add_para('智流工坊采用前后端分离的 B/S 架构，分为四层：')
add_bullet('展示层（前端）：Vue 3 + Vite 单页应用')
add_bullet('应用层（后端）：Spring Boot RESTful API 服务')
add_bullet('引擎层：内置 DAG 任务执行引擎，插件化 Source → Transform → Sink')
add_bullet('基础设施层：MySQL + Redis + Kafka + Flink')

doc.add_heading('3.2 技术栈', level=2)
add_table(['层级', '技术选型', '说明'], [
    ['前端', 'Vue 3 + Vite', '响应式 SPA'],
    ['后端', 'Spring Boot 3 + JPA', 'RESTful API，自动建表'],
    ['数据库', 'MySQL 8.4', '任务元数据、文件存储'],
    ['缓存', 'Redis 7', '会话管理、增量位点'],
    ['消息队列', 'Apache Kafka 3.8', '流数据接入'],
    ['流处理', 'Apache Flink 1.20', '大规模流任务'],
    ['容器化', 'Docker Compose', '一键部署中间件'],
])

doc.add_heading('3.3 项目结构', level=2)
add_table(['目录', '说明'], [
    ['frontend/', 'Vue 3 前端工程'],
    ['backend/', 'Spring Boot 后端工程'],
    ['flink-jobs/', 'Flink 动态管道作业'],
    ['docs/', '项目文档'],
    ['docker-compose.yml', 'Docker 编排文件'],
])
doc.add_page_break()

# === 四 ===
doc.add_heading('四、核心功能详解', level=1)

doc.add_heading('4.1 数据工作区 - 多源数据接入', level=2)
add_para('数据工作区是平台的"数据入口"，管理所有已导入的数据文件。')
add_para('支持的数据格式：', bold=True)
add_bullet('JSON 文件（.json）')
add_bullet('CSV 文件（.csv）')
add_bullet('Excel 文件（.xlsx / .xls）')
add_bullet('XML 文件（.xml）')
add_para('核心能力：', bold=True)
add_bullet('拖拽上传：将文件直接拖入工作区即可完成上传')
add_bullet('自动解析：系统自动识别文件格式，解析为结构化 JSON 数据')
add_bullet('预览功能：上传后可即时预览前 30 条记录')
add_bullet('文件管理：支持重命名、删除、搜索已上传的文件')
add_bullet('记录统计：展示记录数、文件大小、上传时间等')

doc.add_heading('4.2 作业编排 - 可视化 DAG 流程设计器', level=2)
add_para('作业编排是平台核心亮点，提供可视化有向无环图（DAG）流程设计器。')
add_para('界面布局：', bold=True)
add_bullet('左侧面板：组件面板，按 输入/处理/输出/控制 四类分组')
add_bullet('中央画布：拖拽放置节点，连线定义数据流向')
add_bullet('右侧面板：节点配置详情、测试入口')
add_para('输入节点（Source）：', bold=True, size=11, color='2668E8')
add_table(['节点', '功能说明'], [
    ['JSON/CSV/Excel 输入', '读取工作区中已上传的文件'],
    ['MySQL 输入', '从 MySQL 数据库表读取数据'],
    ['JDBC 输入', '从 PostgreSQL/Oracle 等外部数据库读取'],
    ['Redis 输入', '从 Redis List/Set/Hash 读取数据'],
    ['Kafka 输入', '消费 Kafka Topic 中的消息'],
])
add_para('操作亮点：', bold=True)
add_bullet('拖拽添加：从组件面板拖拽节点到画布')
add_bullet('连线编排：从输出端口拖向输入端口定义数据流')
add_bullet('分支支持：一个节点可连接多个下游节点')
add_bullet('节点测试：选中转换节点可即时查看处理结果')
add_bullet('AI 辅助：输入自然语言自动生成流程方案')
add_bullet('只读模式：运行中的任务自动切换为只读')

doc.add_heading('4.3 转换处理插件体系', level=2)
add_para('内置丰富的数据转换插件：')
add_table(['插件', '功能', '典型场景'], [
    ['字段映射', '重命名字段名', '将 customer_name 映射为 客户姓名'],
    ['条件过滤', '按条件筛选数据行', '过滤金额 > 1000 的订单'],
    ['去重', '按指定字段去除重复记录', '按订单号去重'],
    ['字段脱敏', '敏感字段脱敏处理', '手机号中间四位替换为 *'],
    ['字段拼接', '拼接多个字段为新字段', '姓+名 拼接为全名'],
    ['数据补全', '根据数据库补全缺失字段', '从主数据表补全客户地址'],
    ['Redis 查询补全', '按 Redis Hash 补全字段', '用 Redis 缓存补全商品信息'],
    ['数据校验', '校验字段格式并生成审计信息', '校验手机号/邮箱格式'],
    ['类型转换', '转换字段数据类型', '字符串转日期/数字'],
    ['XML/JSON 转换', 'XML 与 JSON 格式互转', 'XML 报文转 JSON 处理'],
])

doc.add_heading('4.4 多种输出目标', level=2)
add_table(['输出目标', '说明'], [
    ['MySQL', '写入 MySQL 动态表，自动建表'],
    ['Kafka', '发送到 Kafka Topic'],
    ['JDBC', '写入 PostgreSQL/Oracle 等外部数据库'],
    ['CSV/Excel', '导出为文件下载'],
    ['Redis', '写入 Redis 缓存'],
])

doc.add_heading('4.5 任务管理与调度', level=2)
add_para('任务管理是平台的"控制中心"。')
add_para('任务生命周期：草稿(DRAFT) → 已发布(PUBLISHED) → 运行中(RUNNING) → 已完成/失败')
add_para('运行模式：', bold=True)
add_bullet('单次运行（ONCE）：执行一次后结束')
add_bullet('定时运行（SCHEDULED）：按 Cron 表达式定时触发')
add_bullet('持续运行（CONTINUOUS）：7x24 小时持续处理')
add_para('管理功能：', bold=True)
add_bullet('创建/编辑/发布/启动/停止/删除任务')
add_bullet('任务搜索与筛选（按创建者、时间范围）')
add_bullet('查看任务详情、DAG 流程图、运行记录')

doc.add_heading('4.6 持续处理与增量模式', level=2)
add_para('支持长期运行任务的完善机制：')
add_bullet('增量消费：基于游标字段（时间戳/自增ID）的增量读取')
add_bullet('处理标记：通过标记字段区分已处理和未处理数据')
add_bullet('Redis 位点：记录消费位点，保证断点续传')
add_bullet('自动重试：失败后自动重试，可配置最大重试次数')

doc.add_heading('4.7 AI 智能助手', level=2)
add_para('集成 AI 智能助手降低使用门槛：')
add_bullet('自然语言描述流程：输入描述，AI 自动生成 DAG 编排方案')
add_bullet('流程建议：分析当前页面上下文，提供优化建议')
add_bullet('图片/文件理解：支持上传截图，AI 自动识别生成流程')
add_bullet('运行分析：对任务运行结果进行智能分析')
add_para('AI 安全策略：仅提供建议，不允许直接操作任务或修改数据。')

doc.add_heading('4.8 Flink 作业集成', level=2)
add_para('支持将任务提交到 Apache Flink 集群：')
add_bullet('Flink 健康检查：实时显示连接状态')
add_bullet('作业提交：转换为 Flink Pipeline 配置并提交')
add_bullet('作业监控：查看已提交作业状态和提交历史')

doc.add_heading('4.9 用户与权限管理', level=2)
add_table(['角色', '说明'], [
    ['超级管理员(SADMIN)', '全部权限，可管理用户'],
    ['管理员(ADMIN)', '全部权限，可管理用户'],
    ['普通用户(USER)', '可创建和管理自己的任务'],
    ['运维只读(OPS_READONLY)', '只读权限，不可修改任务'],
])
add_para('用户功能：登录/登出、个人信息管理、用户管理（管理员）、操作审计。')

doc.add_heading('4.10 运行日志与监控', level=2)
add_bullet('运行记录列表：所有任务执行历史')
add_bullet('多维度筛选：按名称、时间、状态、触发方式')
add_bullet('告警筛选：只查看包含告警的记录')
add_bullet('运行详情：完整输入输出数据和处理统计')
add_bullet('AI 分析：智能识别数据异常')
doc.add_page_break()

# === 五 ===
doc.add_heading('五、技术栈与部署方案', level=1)
doc.add_heading('5.1 部署方式', level=2)
add_para('推荐 Docker Compose 一键部署：', bold=True)
add_bullet('docker-compose up -d 启动全部中间件')
add_bullet('后端：mvn spring-boot:run 启动')
add_bullet('前端：npm run dev 启动开发服务器')

doc.add_heading('5.2 端口规划', level=2)
add_table(['端口', '服务'], [
    ['5173', '前端开发服务器'],
    ['8080', '后端 API'],
    ['3306/3307', 'MySQL'],
    ['6379', 'Redis'],
    ['9092', 'Kafka Broker'],
    ['8081', 'Flink Dashboard'],
])
doc.add_page_break()

# === 六 ===
doc.add_heading('六、演示流程建议', level=1)
steps = [
    ('第一步：概览页（2分钟）', '展示整体界面风格和导航结构，介绍任务统计。'),
    ('第二步：数据工作区（3分钟）', '拖拽上传 CSV/JSON 文件，展示自动解析和预览功能。'),
    ('第三步：作业编排（5分钟）', '创建新任务进入 DAG 编排，拖拽添加节点、连线、配置过滤条件，演示节点测试和 AI 辅助。'),
    ('第四步：任务管理（3分钟）', '展示任务列表和状态流转，启动任务展示运行结果。'),
    ('第五步：运行日志（2分钟）', '查看运行记录、筛选和详情，演示 AI 运行分析。'),
    ('第六步：Flink 集成（可选，2分钟）', '展示 Flink 健康状态和作业提交。'),
    ('第七步：用户管理（可选，1分钟）', '展示角色权限差异和用户管理功能。'),
]
for title, desc in steps:
    add_para(title, bold=True, size=12, color='2E5D8A')
    add_para(desc)
doc.add_page_break()

# === 七 ===
doc.add_heading('七、常见问题与答疑要点', level=1)
qa = [
    ('Q: 和 Apache NiFi、Airflow 有什么区别？', 'A: 智流工坊更轻量、更聚焦数据处理场景。相比 NiFi 提供更简洁界面和更低门槛；相比 Airflow 专注于数据流可视化编排和实时处理。'),
    ('Q: 支持多大数据量？', 'A: 内置引擎适合中小规模（数万至数十万条）。大规模数据可通过 Flink 集群处理。'),
    ('Q: 如何扩展新数据源或插件？', 'A: 后端插件化架构，新增数据源实现 SourceService 接口，新增转换实现 TransformPlugin 接口即可。'),
    ('Q: 数据安全如何保障？', 'A: Token 认证、四级角色权限、字段脱敏、操作审计日志。'),
    ('Q: 是否支持定时调度？', 'A: 支持单次运行、Cron 定时调度、持续运行三种模式。'),
    ('Q: 能否对接已有数据仓库？', 'A: 可以。通过 JDBC 节点连接 PostgreSQL/Oracle 等主流数据库。'),
]
for q, a in qa:
    add_para(q, bold=True, size=11, color='2E5D8A')
    add_para(a, size=11, space_after=12)
doc.add_page_break()

# === 八 ===
doc.add_heading('八、总结与展望', level=1)
doc.add_heading('8.1 项目总结', level=2)
add_bullet('降低门槛：可视化拖拽编排，零代码完成流程设计')
add_bullet('提升效率：丰富内置插件覆盖常见场景')
add_bullet('保障质量：节点级测试、数据校验、脱敏处理')
add_bullet('灵活扩展：插件化架构支持快速扩展')
add_bullet('智能辅助：AI 助手降低使用门槛')

doc.add_heading('8.2 未来规划', level=2)
add_bullet('更多数据源：HDFS、S3、MongoDB')
add_bullet('更丰富插件：聚合计算、窗口函数、ML 特征工程')
add_bullet('工作流调度：任务依赖、DAG 工作流、失败重试策略')
add_bullet('告警通知：邮件/钉钉/企业微信')
add_bullet('数据血缘：追踪数据流转路径')

doc.add_heading('8.3 结束语', level=2)
add_para('"以上就是 A22 智流工坊的完整介绍。我们希望让数据处理变得像画流程图一样简单直观。感谢大家的聆听，欢迎提问和交流！"')

# 保存
output_path = r'D:\a22\docs\A22智流工坊-项目讲解流程稿.docx'
doc.save(output_path)
print(f'Done: {output_path}')
